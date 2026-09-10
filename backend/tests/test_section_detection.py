import io
import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient

from app.db.models import Resume, ResumeSection, ResumeVersion
from app.services.section_detector import SectionDetector
from app.services.text_extractor import TextExtractor


SAMPLE_RESUME_TEXT = """John Doe
john.doe@example.com | (555) 123-4567 | San Francisco, CA

PROFESSIONAL SUMMARY
Senior Software Engineer with 7+ years of experience building scalable backend microservices,
distributed pipelines, and cloud-native systems with Python, FastAPI, and PostgreSQL.

CORE SKILLS
• Languages: Python, Go, TypeScript, SQL
• Frameworks: FastAPI, Flask, React, Node.js
• Databases: PostgreSQL, Redis, MongoDB
• Cloud & DevOps: Docker, Kubernetes, AWS, CI/CD

WORK EXPERIENCE
Senior Backend Engineer | Tech Corp | 2021 - Present
- Architected asynchronous event-driven microservices serving 10M+ daily requests.
- Optimized PostgreSQL database queries reducing p99 latency by 45%.
- Led a cross-functional team of 6 engineers across Agile sprints.

Software Engineer | StartUp Labs | 2018 - 2021
- Developed RESTful APIs using Python and Flask.
- Integrated third-party payment gateways with Stripe.

EDUCATION
Bachelor of Science in Computer Science
University of California, Berkeley | 2014 - 2018
GPA: 3.8 / 4.0

PROJECTS
OpenSource ATS Matcher
- Created an open-source resume parsing tool utilizing NLP and vector search.
- Received 500+ GitHub stars.

CERTIFICATIONS & LICENSES
AWS Certified Solutions Architect - Associate
HashiCorp Certified Terraform Associate
"""


def test_text_extractor_normalization():
    """Verify Unicode, bullet, and spacing normalization."""
    raw = "John Doe\r\n• Point 1\r\n● Point 2\r\n\xa0\xa0\xa0\n\n\n\nWORK EXPERIENCE\n"
    normalized = TextExtractor.normalize_text(raw)

    assert "\r" not in normalized
    assert "- Point 1" in normalized
    assert "- Point 2" in normalized
    assert "\n\n\n" not in normalized
    assert "WORK EXPERIENCE" in normalized


def test_section_detector_heuristic_headings():
    """Verify section detector recognizes standard and alias headings."""
    detector = SectionDetector()
    sections = detector.detect_sections(SAMPLE_RESUME_TEXT)

    section_types = [s.section_type for s in sections]

    assert "contact" in section_types
    assert "summary" in section_types
    assert "skills" in section_types
    assert "experience" in section_types
    assert "education" in section_types
    assert "projects" in section_types
    assert "certifications" in section_types

    # Check order preservation
    orders = [s.section_order for s in sections]
    assert orders == sorted(orders)

    # Check content preservation
    summary_sec = next(s for s in sections if s.section_type == "summary")
    assert "Senior Software Engineer" in summary_sec.content

    skills_sec = next(s for s in sections if s.section_type == "skills")
    assert "Python, Go, TypeScript" in skills_sec.content

    exp_sec = next(s for s in sections if s.section_type == "experience")
    assert "Tech Corp" in exp_sec.content
    assert "StartUp Labs" in exp_sec.content


def test_section_detector_aliases():
    """Verify section aliases map to canonical section types."""
    detector = SectionDetector()

    text = """Candidate Name
career profile:
Passionate engineer.

EMPLOYMENT HISTORY:
Staff Engineer at Enterprise Ltd.

ACADEMIC QUALIFICATIONS:
Master of Science in Software Engineering.

AREAS OF EXPERTISE:
Machine Learning, PyTorch, Scikit-learn.
"""
    sections = detector.detect_sections(text)
    types = {s.section_type for s in sections}

    assert "summary" in types  # 'career profile' -> summary
    assert "experience" in types  # 'EMPLOYMENT HISTORY' -> experience
    assert "education" in types  # 'ACADEMIC QUALIFICATIONS' -> education
    assert "skills" in types  # 'AREAS OF EXPERTISE' -> skills


def test_section_detector_custom_heading():
    """Verify unknown all-caps or underlined headings become custom sections."""
    detector = SectionDetector()

    text = """John Doe
VOLUNTEERING INITIATIVES
Taught code to underprivileged students.

COMMUNITY CONTRIBUTIONS
Organized regional tech meetups.
"""
    sections = detector.detect_sections(text)
    headings = [s.raw_heading for s in sections]

    assert any("VOLUNTEERING INITIATIVES" in h for h in headings)


@pytest.fixture
def auth_user_a(client: TestClient):
    """Register and login candidate User A, returning auth headers."""
    email = "section.user.a@example.com"
    pw = "SecurePassword123!"
    client.post("/api/v1/auth/register", json={
        "email": email,
        "password": pw,
        "full_name": "Section User A",
    })
    token = client.post("/api/v1/auth/login", json={"email": email, "password": pw}).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def auth_user_b(client: TestClient):
    """Register and login candidate User B, returning auth headers."""
    email = "section.user.b@example.com"
    pw = "SecurePassword123!"
    client.post("/api/v1/auth/register", json={
        "email": email,
        "password": pw,
        "full_name": "Section User B",
    })
    token = client.post("/api/v1/auth/login", json={"email": email, "password": pw}).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def test_api_upload_and_auto_section_detection(client: TestClient, auth_user_a: dict):
    """Verify that uploading a resume automatically parses text and populates ResumeSection records."""
    response = client.post(
        "/api/v1/resumes/upload",
        headers=auth_user_a,
        files={"file": ("candidate_resume.txt", SAMPLE_RESUME_TEXT.encode("utf-8"), "text/plain")},
        data={"name": "Auto Parse Test Resume"},
    )
    assert response.status_code == 201
    data = response.json()
    resume_id = data["id"]
    assert data["parser_status"] == "completed"

    # Fetch detected sections via API
    sec_resp = client.get(f"/api/v1/resumes/{resume_id}/sections", headers=auth_user_a)
    assert sec_resp.status_code == 200
    sec_data = sec_resp.json()

    assert sec_data["total_sections"] >= 5
    types = [s["section_type"] for s in sec_data["sections"]]
    assert "summary" in types
    assert "skills" in types
    assert "experience" in types
    assert "education" in types


def test_api_get_resume_text(client: TestClient, auth_user_a: dict):
    """Verify retrieving normalized extracted text."""
    upload_resp = client.post(
        "/api/v1/resumes/upload",
        headers=auth_user_a,
        files={"file": ("test_text.txt", SAMPLE_RESUME_TEXT.encode("utf-8"), "text/plain")},
    )
    resume_id = upload_resp.json()["id"]

    text_resp = client.get(f"/api/v1/resumes/{resume_id}/text", headers=auth_user_a)
    assert text_resp.status_code == 200
    data = text_resp.json()

    assert data["parser_status"] == "completed"
    assert "PROFESSIONAL SUMMARY" in data["extracted_text"]
    assert "WORK EXPERIENCE" in data["extracted_text"]


def test_api_reparse_resume(client: TestClient, auth_user_a: dict):
    """Verify POST /{id}/parse manually triggers reparsing and section detection."""
    upload_resp = client.post(
        "/api/v1/resumes/upload",
        headers=auth_user_a,
        files={"file": ("reparse.txt", SAMPLE_RESUME_TEXT.encode("utf-8"), "text/plain")},
    )
    resume_id = upload_resp.json()["id"]

    parse_resp = client.post(f"/api/v1/resumes/{resume_id}/parse", headers=auth_user_a)
    assert parse_resp.status_code == 200
    data = parse_resp.json()

    assert data["parser_status"] == "completed"
    assert data["total_sections"] >= 5


def test_section_detection_pdf(client: TestClient, auth_user_a: dict):
    """Verify extracting text and detecting sections from a synthetic PDF resume."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Jane Smith\njane@example.com\n\nPROFESSIONAL SUMMARY\nLead AI Architect with deep learning expertise.\n\nTECHNICAL SKILLS\nPyTorch, TensorFlow, Transformers, Python\n\nEXPERIENCE\nAI Lead at Global Corp\n- Built LLM pipelines.\n\nEDUCATION\nPhD in Computer Science, Stanford")
    pdf_bytes = doc.tobytes()
    doc.close()

    upload_resp = client.post(
        "/api/v1/resumes/upload",
        headers=auth_user_a,
        files={"file": ("jane_resume.pdf", pdf_bytes, "application/pdf")},
    )
    assert upload_resp.status_code == 201
    resume_id = upload_resp.json()["id"]

    sec_resp = client.get(f"/api/v1/resumes/{resume_id}/sections", headers=auth_user_a)
    assert sec_resp.status_code == 200
    sec_data = sec_resp.json()

    types = [s["section_type"] for s in sec_data["sections"]]
    assert "summary" in types
    assert "skills" in types
    assert "experience" in types
    assert "education" in types


def test_section_detection_docx(client: TestClient, auth_user_a: dict):
    """Verify extracting text and detecting sections from a real DOCX resume."""
    doc = docx.Document()
    doc.add_heading("Alex Rivera", level=0)
    doc.add_paragraph("alex.rivera@example.com | New York, NY")
    doc.add_heading("SUMMARY", level=1)
    doc.add_paragraph("DevOps Engineer specializing in Kubernetes and cloud reliability.")
    doc.add_heading("SKILLS", level=1)
    doc.add_paragraph("Docker, Kubernetes, Terraform, AWS, Prometheus")
    doc.add_heading("WORK EXPERIENCE", level=1)
    doc.add_paragraph("DevOps Specialist at CloudScale | 2020 - Present")
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph("BS in Information Technology, NYIT")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    upload_resp = client.post(
        "/api/v1/resumes/upload",
        headers=auth_user_a,
        files={"file": ("alex_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert upload_resp.status_code == 201
    resume_id = upload_resp.json()["id"]

    sec_resp = client.get(f"/api/v1/resumes/{resume_id}/sections", headers=auth_user_a)
    assert sec_resp.status_code == 200
    sec_data = sec_resp.json()

    types = [s["section_type"] for s in sec_data["sections"]]
    assert "summary" in types
    assert "skills" in types
    assert "experience" in types
    assert "education" in types


def test_section_access_ownership_isolation(
    client: TestClient,
    auth_user_a: dict,
    auth_user_b: dict,
):
    """User B cannot access or parse User A's resume sections."""
    upload_resp = client.post(
        "/api/v1/resumes/upload",
        headers=auth_user_a,
        files={"file": ("user_a.txt", SAMPLE_RESUME_TEXT.encode("utf-8"), "text/plain")},
    )
    resume_id = upload_resp.json()["id"]

    # User B attempts to read User A's sections
    forbidden_get = client.get(f"/api/v1/resumes/{resume_id}/sections", headers=auth_user_b)
    assert forbidden_get.status_code == 404

    # User B attempts to trigger parse on User A's resume
    forbidden_parse = client.post(f"/api/v1/resumes/{resume_id}/parse", headers=auth_user_b)
    assert forbidden_parse.status_code == 404

