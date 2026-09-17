import io
import fitz
import docx
import pytest


def create_resume_a_pdf() -> bytes:
    """Resume A: Python, FastAPI, PostgreSQL."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "ALICE PYTHONISTA\nBackend Engineer\nalice@example.com")
    page.insert_text((50, 100), "SUMMARY\nBackend specialist in Python, FastAPI, and PostgreSQL.")
    page.insert_text((50, 150), "SKILLS\nPython, FastAPI, PostgreSQL, Docker, Redis")
    page.insert_text((50, 200), "EXPERIENCE\nSenior Python Developer at DataCorp (2021-Present)\nBuilt high-throughput REST APIs using FastAPI and PostgreSQL.")
    page.insert_text((50, 300), "EDUCATION\nB.S. in Computer Science, Tech University")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_resume_b_docx() -> bytes:
    """Resume B: Java, Spring Boot, MySQL."""
    doc = docx.Document()
    doc.add_heading("BOB JAVA DEVELOPER", level=0)
    doc.add_paragraph("bob@example.com | Enterprise Solutions Architect")
    doc.add_heading("SUMMARY", level=1)
    doc.add_paragraph("Enterprise backend engineer with 8 years building microservices in Java and Spring Boot.")
    doc.add_heading("SKILLS", level=1)
    doc.add_paragraph("Java, Spring Boot, MySQL, Hibernate, Maven, Kafka")
    doc.add_heading("EXPERIENCE", level=1)
    doc.add_paragraph("Lead Java Engineer at FinanceHub (2020-Present)\nArchitected banking pipelines with Spring Boot and MySQL.")
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph("M.S. in Software Engineering, State College")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_resume_c_txt() -> bytes:
    """Resume C: React, TypeScript, Node.js."""
    return (
        "CHARLIE FRONTEND\n"
        "charlie@example.com | Web Developer\n\n"
        "PROFESSIONAL SUMMARY\n"
        "Modern frontend engineer specialized in React, TypeScript, and Node.js.\n\n"
        "TECHNICAL SKILLS\n"
        "React, TypeScript, Node.js, Next.js, Tailwind CSS, HTML5, JavaScript\n\n"
        "WORK EXPERIENCE\n"
        "Frontend Lead at WebCraft (2022-Present)\n"
        "Engineered scalable web applications using React, TypeScript, and Node.js.\n\n"
        "EDUCATION\n"
        "B.A. in Interactive Media, Design Institute\n"
    ).encode("utf-8")


@pytest.fixture
def test_user(client):
    email = "candidate.multi@example.com"
    pw = "SecureMultiPass123!"
    client.post("/api/v1/auth/register", json={
        "email": email,
        "password": pw,
        "full_name": "Multi Candidate",
    })
    token = client.post("/api/v1/auth/login", json={"email": email, "password": pw}).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def other_user(client):
    email = "other.candidate@example.com"
    pw = "OtherSecurePass123!"
    client.post("/api/v1/auth/register", json={
        "email": email,
        "password": pw,
        "full_name": "Other Candidate",
    })
    token = client.post("/api/v1/auth/login", json={"email": email, "password": pw}).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def uploaded_three_resumes(client, test_user):
    """Uploads 3 distinct resumes for test_user and returns their response payloads."""
    # 1. Resume A
    res_a = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("Alice_Python.pdf", create_resume_a_pdf(), "application/pdf")},
        data={"name": "Alice Python Staff Resume"},
        headers=test_user,
    ).json()

    # 2. Resume B
    res_b = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("Bob_Java.docx", create_resume_b_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"name": "Bob Java Enterprise Resume"},
        headers=test_user,
    ).json()

    # 3. Resume C
    res_c = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("Charlie_React.txt", create_resume_c_txt(), "text/plain")},
        data={"name": "Charlie React Frontend Resume"},
        headers=test_user,
    ).json()

    return {"a": res_a, "b": res_b, "c": res_c}


def test_three_resumes_upload_and_recognition(client, test_user, uploaded_three_resumes):
    """
    Verify all three uploaded resumes are recognized with distinct metadata,
    independent skills, sections, and unique IDs.
    """
    data_a = uploaded_three_resumes["a"]
    data_b = uploaded_three_resumes["b"]
    data_c = uploaded_three_resumes["c"]

    id_a = data_a["id"]
    id_b = data_b["id"]
    id_c = data_c["id"]

    # Verify unique IDs
    assert len({id_a, id_b, id_c}) == 3

    # Primary flag: first upload is primary, subsequent are not
    assert data_a["is_primary"] is True
    assert data_b["is_primary"] is False
    assert data_c["is_primary"] is False

    # Counts and parser status
    assert data_a["parser_status"] == "completed"
    assert data_b["parser_status"] == "completed"
    assert data_c["parser_status"] == "completed"

    assert data_a["sections_count"] >= 3
    assert data_b["sections_count"] >= 3
    assert data_c["sections_count"] >= 3

    assert data_a["skills_count"] >= 2
    assert data_b["skills_count"] >= 2
    assert data_c["skills_count"] >= 2

    # Verify listing API returns all 3
    list_res = client.get("/api/v1/resumes", headers=test_user)
    assert list_res.status_code == 200
    items = list_res.json()["items"]
    assert len(items) == 3
    item_ids = {item["id"] for item in items}
    assert item_ids == {id_a, id_b, id_c}

    # Verify skills isolation
    skills_a_resp = client.get(f"/api/v1/resumes/{id_a}/skills", headers=test_user)
    skills_a = [s["skill_name"].lower() for s in skills_a_resp.json()]
    assert any("python" in s for s in skills_a)
    assert any("fastapi" in s for s in skills_a)
    assert not any("java" == s or "spring" in s for s in skills_a)
    assert not any("react" in s for s in skills_a)

    skills_b_resp = client.get(f"/api/v1/resumes/{id_b}/skills", headers=test_user)
    skills_b = [s["skill_name"].lower() for s in skills_b_resp.json()]
    assert any("java" == s or "spring" in s for s in skills_b)
    assert not any("fastapi" in s for s in skills_b)
    assert not any("react" in s for s in skills_b)

    skills_c_resp = client.get(f"/api/v1/resumes/{id_c}/skills", headers=test_user)
    skills_c = [s["skill_name"].lower() for s in skills_c_resp.json()]
    assert any("react" in s or "typescript" in s for s in skills_c)
    assert not any("fastapi" in s for s in skills_c)
    assert not any("spring" in s for s in skills_c)


def test_primary_switch_and_deletion(client, test_user, uploaded_three_resumes):
    """Verify setting primary updates correctly and deletion leaves others intact."""
    data_a = uploaded_three_resumes["a"]
    data_b = uploaded_three_resumes["b"]
    data_c = uploaded_three_resumes["c"]

    # Resume A is currently primary. Set Resume B as primary.
    set_res = client.patch(f"/api/v1/resumes/{data_b['id']}/primary", headers=test_user)
    assert set_res.status_code == 200
    assert set_res.json()["is_primary"] is True

    # Re-fetch list and verify only Resume B is primary
    list_after = client.get("/api/v1/resumes", headers=test_user).json()["items"]
    primaries = [i for i in list_after if i["is_primary"]]
    assert len(primaries) == 1
    assert primaries[0]["id"] == data_b["id"]

    # Delete Resume A
    del_res = client.delete(f"/api/v1/resumes/{data_a['id']}", headers=test_user)
    assert del_res.status_code == 200

    # Verify remaining count is 2 (B and C)
    remaining = client.get("/api/v1/resumes", headers=test_user).json()["items"]
    assert len(remaining) == 2
    rem_ids = {i["id"] for i in remaining}
    assert rem_ids == {data_b["id"], data_c["id"]}


def test_analysis_routing_with_selected_resume(client, test_user, uploaded_three_resumes):
    """Verify that running analysis with a specific resume_id analyzes that exact resume."""
    data_b = uploaded_three_resumes["b"]
    res_b_id = data_b["id"]

    # Run analysis explicitly on Resume B
    analysis_res = client.post(
        "/api/v1/analyses",
        json={"resume_id": res_b_id},
        headers=test_user,
    )
    assert analysis_res.status_code == 201
    analysis_data = analysis_res.json()
    assert analysis_data["resume_id"] == res_b_id
    assert analysis_data["resume_name"] == data_b["name"]
    assert analysis_data["overall_score"] is not None
    assert analysis_data["skills_detected_count"] > 0
    assert analysis_data["ats_score"] is not None


def test_cross_user_isolation(client, test_user, other_user, uploaded_three_resumes):
    """User B must never see User A's resumes."""
    data_a = uploaded_three_resumes["a"]
    id_a = data_a["id"]

    # User B lists resumes -> empty
    res_b_list = client.get("/api/v1/resumes", headers=other_user).json()["items"]
    assert len(res_b_list) == 0

    # User B cannot access User A's resume directly
    get_res = client.get(f"/api/v1/resumes/{id_a}", headers=other_user)
    assert get_res.status_code == 404

    # User B cannot delete User A's resume
    del_res = client.delete(f"/api/v1/resumes/{id_a}", headers=other_user)
    assert del_res.status_code == 404

    # User B cannot run analysis on User A's resume
    an_res = client.post("/api/v1/analyses", json={"resume_id": id_a}, headers=other_user)
    assert an_res.status_code == 404
