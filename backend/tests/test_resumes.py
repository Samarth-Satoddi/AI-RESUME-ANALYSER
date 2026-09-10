import io
import os
import uuid
import fitz
import docx
import pytest

from app.core.config import settings
from app.db.models import Resume, ResumeVersion, User
from app.services.storage_service import StorageService


def create_sample_pdf_bytes() -> bytes:
    """Generate a valid, synthetic in-memory PDF."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Alex Morgan - Senior Backend Engineer")
    page.insert_text((50, 100), "Experience: Python, PostgreSQL, Docker")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx_bytes() -> bytes:
    """Generate a valid, synthetic in-memory DOCX."""
    doc = docx.Document()
    doc.add_heading("Alex Morgan - Resume", level=0)
    doc.add_paragraph("Distributed systems engineer specializing in Python and cloud architecture.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_sample_txt_bytes() -> bytes:
    """Generate a valid UTF-8 text resume."""
    return (
        "ALEX MORGAN\n"
        "alex@example.com | San Francisco, CA\n"
        "SUMMARY: Backend engineer with 6 years experience in FastAPI and PostgreSQL.\n"
    ).encode("utf-8")


@pytest.fixture
def auth_user_a(client):
    """Register and login User A, returning auth token and user headers."""
    email = "userA.resume@example.com"
    pw = "SecurePassword123!"
    client.post("/api/v1/auth/register", json={
        "email": email,
        "password": pw,
        "full_name": "User A",
    })
    token = client.post("/api/v1/auth/login", json={"email": email, "password": pw}).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def auth_user_b(client):
    """Register and login User B, returning auth token and user headers."""
    email = "userB.resume@example.com"
    pw = "SecurePassword456!"
    client.post("/api/v1/auth/register", json={
        "email": email,
        "password": pw,
        "full_name": "User B",
    })
    token = client.post("/api/v1/auth/login", json={"email": email, "password": pw}).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


# ==========================================
# 1. FILE VALIDATION & FORMAT TESTS
# ==========================================

def test_upload_valid_pdf(client, auth_user_a):
    """Test successful upload of a genuine PDF document."""
    pdf_bytes = create_sample_pdf_bytes()
    files = {"file": ("alex_resume.pdf", pdf_bytes, "application/pdf")}
    data = {"name": "Alex Morgan Staff Resume"}

    response = client.post("/api/v1/resumes/upload", files=files, data=data, headers=auth_user_a)
    assert response.status_code == 201
    res = response.json()
    assert res["name"] == "Alex Morgan Staff Resume"
    assert res["original_filename"] == "alex_resume.pdf"
    assert res["file_type"] == "pdf"
    assert res["file_size"] == len(pdf_bytes)
    assert res["is_primary"] is True  # First upload becomes primary
    assert res["version_number"] == 1
    assert res["parser_status"] in ("pending", "completed")



def test_upload_valid_docx(client, auth_user_a):
    """Test successful upload of a genuine DOCX document and verify primary flag on multiple uploads."""
    # First upload (becomes primary)
    pdf_bytes = create_sample_pdf_bytes()
    r1 = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("first.pdf", pdf_bytes, "application/pdf")},
        headers=auth_user_a,
    ).json()
    assert r1["is_primary"] is True

    # Second upload (DOCX - should not be primary)
    docx_bytes = create_sample_docx_bytes()
    files = {"file": ("alex_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    response = client.post("/api/v1/resumes/upload", files=files, headers=auth_user_a)
    assert response.status_code == 201
    res = response.json()
    assert res["file_type"] == "docx"
    assert res["is_primary"] is False  # Second upload is not primary by default


def test_upload_valid_txt(client, auth_user_a):
    """Test successful upload of a genuine TXT document."""
    txt_bytes = create_sample_txt_bytes()
    files = {"file": ("alex_resume.txt", txt_bytes, "text/plain")}

    response = client.post("/api/v1/resumes/upload", files=files, headers=auth_user_a)
    assert response.status_code == 201
    res = response.json()
    assert res["file_type"] == "txt"


def test_upload_unsupported_extension_rejected(client, auth_user_a):
    """Test that non-resume extensions (.exe, .zip, .png, etc.) are rejected."""
    files = {"file": ("malicious.exe", b"MZexecutabledata", "application/octet-stream")}
    response = client.post("/api/v1/resumes/upload", files=files, headers=auth_user_a)
    assert response.status_code == 422
    assert "Unsupported file format" in response.json()["message"]


def test_upload_empty_file_rejected(client, auth_user_a):
    """Test that 0-byte files are rejected."""
    files = {"file": ("empty.pdf", b"", "application/pdf")}
    response = client.post("/api/v1/resumes/upload", files=files, headers=auth_user_a)
    assert response.status_code == 422
    assert "empty" in response.json()["message"]


def test_upload_corrupted_pdf_rejected(client, auth_user_a):
    """Test that malformed PDF content (missing %PDF- or corrupt body) is rejected."""
    corrupt_bytes = b"not a real pdf content just random bytes"
    files = {"file": ("corrupt.pdf", corrupt_bytes, "application/pdf")}
    response = client.post("/api/v1/resumes/upload", files=files, headers=auth_user_a)
    assert response.status_code == 422
    assert "Invalid PDF" in response.json()["message"]


def test_upload_corrupted_docx_rejected(client, auth_user_a):
    """Test that invalid zip archive masquerading as .docx is rejected."""
    corrupt_docx = b"PK\x03\x04fakearchivedatawithoutproperxml"
    files = {"file": ("corrupt.docx", corrupt_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/v1/resumes/upload", files=files, headers=auth_user_a)
    assert response.status_code == 422


def test_upload_binary_masquerading_as_txt_rejected(client, auth_user_a):
    """Test that binary/executable content with .txt extension is rejected."""
    binary_bytes = b"MZ\x90\x00\x03\x00\x00\x00\x04\x00\x00\x00\xff\xff\x00\x00"
    files = {"file": ("binary.txt", binary_bytes, "text/plain")}
    response = client.post("/api/v1/resumes/upload", files=files, headers=auth_user_a)
    assert response.status_code == 422


def test_path_traversal_sanitization(client, auth_user_a):
    """Test that path traversal attempts in filename (../../etc/passwd) are safely sanitized."""
    pdf_bytes = create_sample_pdf_bytes()
    files = {"file": ("../../evil_payload.pdf", pdf_bytes, "application/pdf")}
    response = client.post("/api/v1/resumes/upload", files=files, headers=auth_user_a)
    assert response.status_code == 201
    res = response.json()
    assert ".." not in res["original_filename"]
    assert "/" not in res["original_filename"]
    assert "\\" not in res["original_filename"]
    assert "evil_payload.pdf" in res["original_filename"]


# ==========================================
# 2. AUTHENTICATION & OWNERSHIP TESTS
# ==========================================

def test_upload_unauthenticated_rejected(client):
    """Test that unauthenticated upload returns HTTP 401."""
    pdf_bytes = create_sample_pdf_bytes()
    files = {"file": ("resume.pdf", pdf_bytes, "application/pdf")}
    response = client.post("/api/v1/resumes/upload", files=files)
    assert response.status_code == 401


def test_user_ownership_isolation(client, auth_user_a, auth_user_b):
    """Test that User B cannot read, download, or delete User A's resume."""
    # User A uploads a resume
    pdf_bytes = create_sample_pdf_bytes()
    files = {"file": ("private_a.pdf", pdf_bytes, "application/pdf")}
    create_res = client.post("/api/v1/resumes/upload", files=files, headers=auth_user_a)
    assert create_res.status_code == 201
    resume_id = create_res.json()["id"]

    # User B attempts to read User A's resume -> 404
    r_get = client.get(f"/api/v1/resumes/{resume_id}", headers=auth_user_b)
    assert r_get.status_code == 404

    # User B attempts to download User A's resume -> 404
    r_dl = client.get(f"/api/v1/resumes/{resume_id}/download", headers=auth_user_b)
    assert r_dl.status_code == 404

    # User B attempts to delete User A's resume -> 404
    r_del = client.delete(f"/api/v1/resumes/{resume_id}", headers=auth_user_b)
    assert r_del.status_code == 404

    # User A can read and download their own resume
    assert client.get(f"/api/v1/resumes/{resume_id}", headers=auth_user_a).status_code == 200
    assert client.get(f"/api/v1/resumes/{resume_id}/download", headers=auth_user_a).status_code == 200


# ==========================================
# 3. LIFECYCLE, PRIMARY TOGGLE & CLEANUP TESTS
# ==========================================

def test_set_primary_resume(client, auth_user_a):
    """Test atomic primary resume toggle."""
    # Upload 2 resumes
    pdf_bytes = create_sample_pdf_bytes()
    r1 = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("r1.pdf", pdf_bytes, "application/pdf")},
        headers=auth_user_a,
    ).json()
    r2 = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("r2.pdf", pdf_bytes, "application/pdf")},
        headers=auth_user_a,
    ).json()

    # Set r2 as primary
    res = client.patch(f"/api/v1/resumes/{r2['id']}/primary", headers=auth_user_a)
    assert res.status_code == 200
    assert res.json()["is_primary"] is True

    # Verify r1 is now not primary
    r1_check = client.get(f"/api/v1/resumes/{r1['id']}", headers=auth_user_a).json()
    assert r1_check["is_primary"] is False


def test_delete_resume_and_file_cleanup(client, auth_user_a, db_session):
    """Test deleting resume removes DB record and clears stored file from disk."""
    pdf_bytes = create_sample_pdf_bytes()
    upload_res = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("to_delete.pdf", pdf_bytes, "application/pdf")},
        headers=auth_user_a,
    ).json()
    resume_id = upload_res["id"]

    # Verify stored file exists on disk
    resume_orm = db_session.query(Resume).filter_by(id=uuid.UUID(resume_id)).first()
    assert resume_orm is not None
    stored_path = StorageService.get_absolute_path(resume_orm.storage_path)
    assert os.path.exists(stored_path)

    # Delete resume via API
    del_res = client.delete(f"/api/v1/resumes/{resume_id}", headers=auth_user_a)
    assert del_res.status_code == 200

    # Verify file is deleted from disk
    assert not os.path.exists(stored_path)
