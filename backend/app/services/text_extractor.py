import os
import re
import unicodedata
from pathlib import Path
from typing import Optional
import pymupdf
import docx
from app.core.exceptions import ValidationException


class TextExtractor:
    """Service to extract and normalize text from PDF, DOCX, and TXT resumes."""

    @staticmethod
    def extract_from_pdf(file_path: Path) -> str:
        """Extract plain text from a PDF document using PyMuPDF."""
        try:
            doc = pymupdf.open(file_path)
            extracted_pages = []
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text("text")
                if text.strip():
                    extracted_pages.append(text.strip())
            doc.close()
            return "\n\n".join(extracted_pages)
        except Exception as e:
            raise ValidationException(f"Failed to extract text from PDF document: {str(e)}")

    @staticmethod
    def extract_from_docx(file_path: Path) -> str:
        """Extract text from a DOCX document including paragraphs and tables."""
        try:
            doc = docx.Document(file_path)
            elements = []

            # Extract body paragraphs
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    elements.append(text)

            # Also extract text from any tables (many resumes format skills/columns in tables)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        # Deduplicate identical adjacent cells (due to merged cells)
                        deduped = []
                        for cell in row_cells:
                            if not deduped or deduped[-1] != cell:
                                deduped.append(cell)
                        elements.append(" | ".join(deduped))

            return "\n\n".join(elements)
        except Exception as e:
            raise ValidationException(f"Failed to extract text from DOCX document: {str(e)}")

    @staticmethod
    def extract_from_txt(file_path: Path) -> str:
        """Extract plain text from a TXT document handling UTF-8 and Latin-1."""
        try:
            with open(file_path, "rb") as f:
                content = f.read()

            try:
                return content.decode("utf-8")
            except UnicodeDecodeError:
                return content.decode("latin-1", errors="replace")
        except Exception as e:
            raise ValidationException(f"Failed to extract text from TXT document: {str(e)}")

    @classmethod
    def normalize_text(cls, raw_text: str) -> str:
        """
        Normalize extracted text:
        - Unicode normalization (NFKC)
        - Bullet point standardization
        - Carriage return standardization
        - Excess blank line collapse
        - Preservation of structural line breaks for section headers
        """
        if not raw_text:
            return ""

        # 1. Unicode NFKC normalization
        normalized = unicodedata.normalize("NFKC", raw_text)

        # 2. Standardize linebreaks to \n
        normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")

        # 3. Standardize non-breaking spaces and tabs
        normalized = normalized.replace("\xa0", " ").replace("\t", "    ")

        # 4. Standardize bullet characters to standard '- '
        bullet_pattern = r"^[\s]*[•●▪▫◦\*\u2022\u2023\u25E6\u2043\u2219]\s*"
        lines = []
        for line in normalized.split("\n"):
            stripped_line = line.rstrip()
            # Replace bullet markers at beginning of line
            subbed = re.sub(bullet_pattern, "- ", stripped_line)
            lines.append(subbed)

        # 5. Join lines and collapse 3+ consecutive newlines to 2
        normalized = "\n".join(lines)
        normalized = re.sub(r"\n{3,}", "\n\n", normalized)

        return normalized.strip()

    @classmethod
    def extract(cls, file_path: Path, file_type: str) -> str:
        """Extract and normalize text according to file extension."""
        ft = file_type.lower().strip(".").strip()
        if ft == "pdf":
            raw = cls.extract_from_pdf(file_path)
        elif ft in ("docx", "doc"):
            raw = cls.extract_from_docx(file_path)
        elif ft == "txt":
            raw = cls.extract_from_txt(file_path)
        else:
            raise ValidationException(f"Unsupported file format for text extraction: {file_type}")

        return cls.normalize_text(raw)
